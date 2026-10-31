"""Offline end-to-end search demo (no AWS, no Docker, no server).

Seeds the local hybrid index (embedded Qdrant + SQLite FTS5) from a few bundled
sample compliance documents using the deterministic fake embedder, then runs
queries and prints ranked passages — exactly what the UI shows.

It writes to ./data/search and releases its locks on exit, so afterwards you can
browse the SAME index in the web UI:

    # PowerShell
    $env:SEARCH_EMBED_PROVIDER = "fake"
    uv run python scripts/search_demo.py            # seed + run demo queries
    uv run uvicorn rag_search.app:app --port 8004   # open http://localhost:8004

Usage:
    uv run python scripts/search_demo.py                       # default queries
    uv run python scripts/search_demo.py "encryption at rest"  # custom query
"""

from __future__ import annotations

import asyncio
import sys
from collections.abc import AsyncIterator
from io import BytesIO
from typing import Any

from rag_search.chunking import Chunker
from rag_search.config import SearchSettings
from rag_search.embedder import FakeEmbedder
from rag_search.indexer import IndexingService
from rag_search.models import IndexRequest
from rag_search.retrieval import SearchService
from rag_search.stores import KeywordIndex, LocalVectorIndex

from rag_core.aws import S3ObjectSummary
from rag_core.config import S3Settings, Settings, SQSSettings
from rag_core.logging import configure_logging
from rag_processing.config import ProcessingSettings
from rag_processing.extractors import build_default_registry
from rag_processing.ocr.noop import NoOpOCREngine
from rag_processing.service import ProcessingService
from rag_processing.store import InMemoryContentStore

DATA_DIR = "./data/search"


def _docx(title: str, body: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=1)
    for para in body.split("\n"):
        doc.add_paragraph(para)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _html(title: str, body: str) -> bytes:
    paras = "".join(f"<p>{p}</p>" for p in body.split("\n"))
    return f"<html><body><h1>{title}</h1>{paras}</body></html>".encode()


SAMPLES: dict[str, bytes] = {
    "foxtrot/information_security_policy.docx": _docx(
        "Information Security Policy",
        "All confidential data must be encrypted at rest using AES-256.\n"
        "Data in transit must use TLS 1.2 or higher.\n"
        "Access requires multi-factor authentication for all production systems.",
    ),
    "foxtrot/supplier_register.html": _html(
        "Approved Supplier Register",
        "Supplier ACME Corporation supplies raw polymer materials.\n"
        "ACME is located in Germany and holds ISO 9001 and ISO 14001 certification.\n"
        "Supplier Globex provides packaging and is based in France.",
    ),
    "foxtrot/quality_manual.txt": (
        b"Quality Manual. This manual describes the change-control procedure. "
        b"Every design change requires a documented risk assessment and approval "
        b"by the quality manager before release."
    ),
}


class _FakeS3:
    def __init__(self, objects: dict[str, bytes]) -> None:
        self._objects = objects

    async def list_objects(
        self, bucket: str, prefix: str = "", *, page_size: int = 1000
    ) -> AsyncIterator[S3ObjectSummary]:
        for key in self._objects:
            if key.startswith(prefix):
                yield S3ObjectSummary(bucket=bucket, key=key, size_bytes=len(self._objects[key]), etag="e")

    async def get_bytes(self, bucket: str, key: str, *, version_id: str | None = None) -> bytes:
        return self._objects[key]

    async def put_bytes(self, *a: Any, **k: Any) -> None:
        return None

    async def head_object(self, *a: Any, **k: Any) -> dict[str, Any]:
        return {}


class _NullSQS:
    async def send(self, *a: Any, **k: Any) -> str:
        return "noop"


async def _run(queries: list[str]) -> None:
    core = Settings(s3=S3Settings(bucket="demo-bucket", prefix="foxtrot/"), sqs=SQSSettings())
    # Use the same collection the web app defaults to, so the seeded index is the
    # one the UI reads (run the UI with SEARCH_EMBED_PROVIDER=fake to match dims).
    ss = SearchSettings(data_dir=DATA_DIR)

    s3 = _FakeS3(SAMPLES)
    embedder = FakeEmbedder(dim=256)
    vector = LocalVectorIndex(f"{DATA_DIR}/qdrant", collection=ss.collection, dim=256)
    keyword = KeywordIndex(f"{DATA_DIR}/keyword.db")
    content_store = InMemoryContentStore()
    processing = ProcessingService(
        s3=s3, sqs=_NullSQS(), registry=build_default_registry(), ocr_engine=NoOpOCREngine(),
        content_store=content_store, core_settings=core,
        settings=ProcessingSettings(run_worker=False, output_queue_url=""),
    )
    indexer = IndexingService(
        s3=s3, processing=processing, content_store=content_store, chunker=Chunker(),
        embedder=embedder, vector=vector, keyword=keyword, core_settings=core, settings=ss,
    )
    search = SearchService(embedder=embedder, vector=vector, keyword=keyword, settings=ss)

    try:
        summary = await indexer.index(IndexRequest(reset=True))
        print(f"\nIndexed {summary.documents_indexed} documents · {summary.chunks_indexed} chunks "
              f"in {summary.duration_s}s  (data in {DATA_DIR})\n")

        for q in queries:
            resp = await search.search(q, top_k=3)
            print("=" * 78)
            print(f"QUERY: {q}   ({resp.took_ms} ms)")
            print("-" * 78)
            for i, r in enumerate(resp.results, 1):
                loc = f"{r.section or ''} p.{r.page_start}".strip()
                print(f"  {i}. [{r.matched_by:7}] score={r.score:.4f}  {r.source_key}  {loc}")
                print(f"     {r.text[:150].strip()}...")
            if not resp.results:
                print("  (no matches)")
            print()
    finally:
        await vector.aclose()
        await keyword.aclose()


def main() -> None:
    configure_logging(level="WARNING", json_logs=False, service_name="search-demo")
    queries = sys.argv[1:] or [
        "how must confidential data be encrypted at rest",
        "which suppliers are in Germany with ISO certification",
        "what is required before a design change is released",
    ]
    asyncio.run(_run(queries))


if __name__ == "__main__":
    main()
