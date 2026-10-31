"""Local demo of the write path (ingest -> extract), runnable with zero infra.

This is the closest thing to "retrieve the correct document" available today:
it fetches a document and shows the *faithfully extracted* content (pages,
headings, tables, language, classification, PII). Semantic search arrives with
the chunking/embedding/indexer/retrieval milestones (#3-#9).

Usage
-----
# No infra at all — generates a sample DOCX in memory and extracts it:
    uv run python scripts/demo_pipeline.py process --source fake

# Against a real S3 object (needs AWS creds + s3:GetObject on the bucket):
    uv run python scripts/demo_pipeline.py process --source s3 \
        --bucket example-index-bucket --key Test/your-file.pdf

# Dry-run discovery over an S3 prefix (lists new/changed/dup; emits nothing):
    uv run python scripts/demo_pipeline.py list --prefix Test/
"""

from __future__ import annotations

import argparse
import asyncio
from io import BytesIO
from typing import Any

from rag_core.config import get_settings
from rag_core.di import Container
from rag_core.logging import configure_logging
from rag_core.utils import document_id, infer_document_type
from rag_ingestion.config import IngestionSettings
from rag_ingestion.dedup import InMemoryDedupLedger
from rag_ingestion.models import SyncRequest
from rag_ingestion.service import IngestionService
from rag_processing.config import OCRProvider, ProcessingSettings
from rag_processing.extractors import build_default_registry
from rag_processing.ocr import build_ocr_engine
from rag_processing.service import ProcessingService
from rag_processing.store import InMemoryContentStore


# --------------------------------------------------------------------------- #
# Minimal in-memory stand-ins so the "fake" demo needs no AWS at all.
# --------------------------------------------------------------------------- #
class _MemS3:
    def __init__(self, objects: dict[str, bytes]) -> None:
        self._objects = objects

    async def get_bytes(self, bucket: str, key: str, *, version_id: str | None = None) -> bytes:
        return self._objects[key]

    async def put_bytes(self, *args: Any, **kwargs: Any) -> None:
        return None

    async def head_object(self, *args: Any, **kwargs: Any) -> dict[str, Any]:
        return {}


class _NullSQS:
    async def send(self, *args: Any, **kwargs: Any) -> str:
        return "noop"


def _sample_docx() -> bytes:
    """A small, realistic compliance-style DOCX generated in memory."""
    from docx import Document

    doc = Document()
    doc.add_heading("Information Security Policy", level=1)
    doc.add_paragraph(
        "This policy defines acceptable use of company systems. For access, "
        "contact security@example.com. It applies to all employees."
    )
    doc.add_heading("2. Data Handling", level=2)
    doc.add_paragraph("Confidential data must be encrypted at rest and in transit.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Classification"
    table.cell(0, 1).text = "Control"
    table.cell(1, 0).text = "Confidential"
    table.cell(1, 1).text = "AES-256"
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _print_result(result: Any, content: Any) -> None:
    print("\n" + "=" * 70)
    print(f"  document_id   : {result.document_id}")
    print(f"  type          : {result.doc_type.value}")
    print(f"  pages         : {result.page_count}")
    print(f"  tables        : {result.table_count}")
    print(f"  language      : {result.language.value}")
    print(f"  classification: {result.document_class.value}")
    print(f"  PII findings  : {result.pii_findings}")
    print(f"  has OCR       : {result.has_ocr}")
    print(f"  chars         : {result.chars_extracted}")
    print(f"  content URI   : {result.content_uri}")
    print(f"  duration      : {result.duration_s}s")
    if content is not None:
        headings = [
            b.text for p in content.pages for b in p.blocks if b.role == "heading"
        ]
        print(f"\n  headings detected: {headings}")
        snippet = content.full_text[:300].replace("\n", " ")
        print(f"\n  text snippet:\n    {snippet}...")
        for page in content.pages:
            for tbl in page.tables:
                print(f"\n  table (p.{tbl.page}):\n{tbl.markdown}")
    print("=" * 70 + "\n")


async def _process(args: argparse.Namespace) -> None:
    settings = get_settings()
    ps = ProcessingSettings(run_worker=False, output_queue_url="")
    registry = build_default_registry()
    store = InMemoryContentStore()

    if args.source == "fake":
        key = "Test/sample_policy.docx"
        s3: Any = _MemS3({key: _sample_docx()})
        sqs: Any = _NullSQS()
        ocr: Any = build_ocr_engine(
            ProcessingSettings(ocr_provider=OCRProvider.NONE), Container(settings).aws_factory()
        )
        bucket = settings.s3.bucket
    else:
        if not args.key:
            raise SystemExit("--key is required for --source s3")
        key = args.key
        bucket = args.bucket or settings.s3.bucket
        container = Container(settings)
        s3 = container.s3()
        sqs = container.sqs()
        ocr = build_ocr_engine(ps, container.aws_factory())

    service = ProcessingService(
        s3=s3, sqs=sqs, registry=registry, ocr_engine=ocr,
        content_store=store, core_settings=settings, settings=ps,
    )
    doc_id = document_id(bucket, key)
    result = await service.process_object(
        bucket=bucket, key=key, document_id=doc_id, doc_type=infer_document_type(key)
    )
    content = await store.get(doc_id, namespace="default")
    _print_result(result, content)

    if args.source == "s3":
        await ocr.aclose()
        await Container(settings).aclose()


async def _list(args: argparse.Namespace) -> None:
    settings = get_settings()
    container = Container(settings)
    service = IngestionService(
        s3=container.s3(),
        sqs=container.sqs(),
        ledger=InMemoryDedupLedger(),
        core_settings=settings,
        settings=IngestionSettings(emit_to_sqs=False, change_detection="etag"),
    )
    summary = await service.sync(SyncRequest(prefix=args.prefix, dry_run=True))
    print("\n" + "=" * 60)
    print(f"  bucket   : {summary.bucket}")
    print(f"  prefix   : {summary.prefix}")
    print(f"  scanned  : {summary.scanned}")
    print(f"  new      : {summary.new}")
    print(f"  changed  : {summary.changed}")
    print(f"  duplicate: {summary.duplicate}")
    print(f"  skipped  : {summary.skipped}")
    print(f"  errors   : {summary.errors}")
    print(f"  duration : {summary.duration_s}s")
    print("=" * 60 + "\n")
    await container.aclose()


def main() -> None:
    configure_logging(level="WARNING", json_logs=False, service_name="demo")
    parser = argparse.ArgumentParser(description="Agentic RAG write-path demo")
    sub = parser.add_subparsers(dest="command", required=True)

    p = sub.add_parser("process", help="extract one document")
    p.add_argument("--source", choices=["fake", "s3"], default="fake")
    p.add_argument("--bucket", default=None)
    p.add_argument("--key", default=None)

    ls = sub.add_parser("list", help="dry-run discovery over an S3 prefix")
    ls.add_argument("--prefix", default=None)

    args = parser.parse_args()
    if args.command == "process":
        asyncio.run(_process(args))
    elif args.command == "list":
        asyncio.run(_list(args))


if __name__ == "__main__":
    main()
