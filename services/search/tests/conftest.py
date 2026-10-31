"""Fixtures + fakes for search tests (fully offline: fake S3, fake embedder)."""

from __future__ import annotations

from collections.abc import AsyncIterator, Callable
from io import BytesIO

import pytest
from rag_search.chunking import Chunker
from rag_search.config import SearchSettings
from rag_search.embedder import FakeEmbedder
from rag_search.indexer import IndexingService
from rag_search.metadata_extract import MetadataExtractor
from rag_search.metadata_store import MetadataStore
from rag_search.retrieval import SearchService
from rag_search.stores import KeywordIndex, LocalVectorIndex

from rag_core.aws import S3ObjectSummary
from rag_core.config import S3Settings, Settings, SQSSettings
from rag_processing.config import ProcessingSettings
from rag_processing.extractors import build_default_registry
from rag_processing.ocr.noop import NoOpOCREngine
from rag_processing.service import ProcessingService
from rag_processing.store import InMemoryContentStore


class FakeS3:
    def __init__(self, objects: dict[str, bytes]) -> None:
        self._objects = objects

    async def list_objects(
        self, bucket: str, prefix: str = "", *, page_size: int = 1000
    ) -> AsyncIterator[S3ObjectSummary]:
        for key in self._objects:
            if key.startswith(prefix):
                yield S3ObjectSummary(
                    bucket=bucket, key=key, size_bytes=len(self._objects[key]), etag="e"
                )

    async def get_bytes(self, bucket: str, key: str, *, version_id: str | None = None) -> bytes:
        return self._objects[key]

    async def put_bytes(self, *args: object, **kwargs: object) -> None:
        return None

    async def head_object(self, *args: object, **kwargs: object) -> dict[str, object]:
        return {}


class _NullSQS:
    async def send(self, *args: object, **kwargs: object) -> str:
        return "noop"


def make_docx(title: str, body: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=1)
    for para in body.split("\n"):
        doc.add_paragraph(para)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_html(title: str, body: str) -> bytes:
    paras = "".join(f"<p>{p}</p>" for p in body.split("\n"))
    return f"<html><body><h1>{title}</h1>{paras}</body></html>".encode()


@pytest.fixture
def sample_objects() -> dict[str, bytes]:
    return {
        "foxtrot/security_policy.docx": make_docx(
            "Information Security Policy",
            "All confidential data must be encrypted at rest using AES-256 encryption.\n"
            "Access control requires multi-factor authentication for all systems.",
        ),
        "foxtrot/supplier_register.html": make_html(
            "Supplier Register",
            "Supplier ACME Corporation provides raw materials for production.\n"
            "ACME is located in Germany and holds ISO 9001 certification.",
        ),
    }


@pytest.fixture
def core_settings() -> Settings:
    return Settings(
        s3=S3Settings(bucket="demo-bucket", prefix="foxtrot/"),
        sqs=SQSSettings(),
    )


@pytest.fixture
def search_settings(tmp_path) -> SearchSettings:
    return SearchSettings(
        data_dir=str(tmp_path / "search"),
        collection="test",
        child_words=40,
        child_overlap=10,
        parent_max_words=120,
        candidate_k=20,
        top_k=5,
    )


@pytest.fixture
def build_stack(
    core_settings: Settings, search_settings: SearchSettings
) -> Callable[..., tuple[IndexingService, SearchService, MetadataStore, list]]:
    def _build(
        objects: dict[str, bytes],
        *,
        mongo: object | None = None,
        stores: tuple[object, object, object] | None = None,
    ) -> tuple[IndexingService, SearchService, MetadataStore, list]:
        """``stores`` injects an alternative (vector, keyword, metadata) triple — the
        Postgres backend reuses this to run the identical pipeline for parity tests.
        Injected stores are NOT returned as closeables; their owner closes them."""
        s3 = FakeS3(objects)
        embedder = FakeEmbedder(dim=256)
        if stores is not None:
            vector, keyword, injected_metadata = stores
        else:
            vector = LocalVectorIndex(
                search_settings.data_dir + "/qdrant",
                collection=search_settings.collection,
                dim=256,
            )
            keyword = KeywordIndex(search_settings.data_dir + "/kw.db")
            injected_metadata = None
        content_store = InMemoryContentStore()
        processing = ProcessingService(
            s3=s3,  # type: ignore[arg-type]
            sqs=_NullSQS(),  # type: ignore[arg-type]
            registry=build_default_registry(),
            ocr_engine=NoOpOCREngine(),
            content_store=content_store,
            core_settings=core_settings,
            settings=ProcessingSettings(run_worker=False, output_queue_url=""),
        )
        chunker = Chunker(parent_max_words=120, child_words=40, child_overlap=10)
        metadata_store = injected_metadata or MetadataStore(
            search_settings.data_dir + "/meta.db"
        )
        # Heuristic-only metadata (no LLM) keeps indexing offline in tests.
        metadata_extractor = MetadataExtractor(None, use_llm=False)
        indexer = IndexingService(
            s3=s3,  # type: ignore[arg-type]
            processing=processing,
            content_store=content_store,
            chunker=chunker,
            embedder=embedder,
            vector=vector,
            keyword=keyword,
            metadata_extractor=metadata_extractor,
            metadata_store=metadata_store,
            core_settings=core_settings,
            settings=search_settings,
            mongo=mongo,  # type: ignore[arg-type]
        )
        search = SearchService(
            embedder=embedder, vector=vector, keyword=keyword, settings=search_settings
        )
        closeables = [] if stores is not None else [vector, keyword, metadata_store]
        return indexer, search, metadata_store, closeables

    return _build
