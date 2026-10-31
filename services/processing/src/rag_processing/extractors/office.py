"""Microsoft Office extractors: DOCX, PPTX, XLSX.

* DOCX: walks body content in document order (paragraphs + tables interleaved),
  maps Word heading styles to heading blocks, paginates flat flow into synthetic
  pages.
* PPTX: one slide = one page; titles become headings, tables preserved.
* XLSX: one worksheet = one page, rendered as a table (data-only values).
"""

from __future__ import annotations

import asyncio
from io import BytesIO

from rag_core.models.content import BlockType, ExtractedContent, Page, Table, TextBlock
from rag_core.models.enums import DocumentType
from rag_processing.extractors.base import ExtractionContext, Extractor
from rag_processing.extractors.layout import grid_to_markdown, paginate_items

_XLSX_MAX_ROWS = 5000
_XLSX_MAX_COLS = 100


class DocxExtractor(Extractor):
    supported = (DocumentType.DOCX,)

    async def extract(self, data: bytes, ctx: ExtractionContext) -> ExtractedContent:
        items = await asyncio.to_thread(self._parse, data)
        return ExtractedContent(document_id=ctx.document_id, pages=paginate_items(items))

    @staticmethod
    def _parse(data: bytes) -> list[TextBlock | Table]:
        from docx import Document
        from docx.table import Table as DocxTable
        from docx.text.paragraph import Paragraph

        doc = Document(BytesIO(data))
        items: list[TextBlock | Table] = []
        # iter_inner_content preserves paragraph/table ordering (python-docx >=1.1).
        for block in doc.iter_inner_content():
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if not text:
                    continue
                style = (block.style.name if block.style else "") or ""
                level = _heading_level_from_style(style)
                items.append(
                    TextBlock(
                        text=text,
                        role=BlockType.HEADING if level else BlockType.PARAGRAPH,
                        page=1,
                        heading_level=level,
                    )
                )
            elif isinstance(block, DocxTable):
                grid = [[cell.text.strip() for cell in row.cells] for row in block.rows]
                grid = [r for r in grid if any(c for c in r)]
                if grid:
                    items.append(Table(page=1, rows=grid, markdown=grid_to_markdown(grid)))
        return items


class PptxExtractor(Extractor):
    supported = (DocumentType.PPTX,)

    async def extract(self, data: bytes, ctx: ExtractionContext) -> ExtractedContent:
        pages = await asyncio.to_thread(self._parse, data, ctx.document_id)
        return ExtractedContent(document_id=ctx.document_id, pages=pages)

    @staticmethod
    def _parse(data: bytes, document_id: str) -> list[Page]:
        from pptx import Presentation

        prs = Presentation(BytesIO(data))
        pages: list[Page] = []
        for idx, slide in enumerate(prs.slides, start=1):
            blocks: list[TextBlock] = []
            tables: list[Table] = []
            title_text = slide.shapes.title.text.strip() if slide.shapes.title else ""
            if title_text:
                blocks.append(
                    TextBlock(text=title_text, role=BlockType.HEADING, page=idx, heading_level=1)
                )
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = "".join(run.text for run in para.runs).strip()
                        if text and text != title_text:
                            blocks.append(TextBlock(text=text, role=BlockType.PARAGRAPH, page=idx))
                elif shape.has_table:
                    tbl = shape.table
                    grid = [[cell.text.strip() for cell in row.cells] for row in tbl.rows]
                    grid = [r for r in grid if any(c for c in r)]
                    if grid:
                        tables.append(Table(page=idx, rows=grid, markdown=grid_to_markdown(grid)))
            text = "\n".join(b.text for b in blocks)
            pages.append(Page(number=idx, text=text, blocks=blocks, tables=tables))
        return pages or [Page(number=1, text="")]


class XlsxExtractor(Extractor):
    supported = (DocumentType.XLSX,)

    async def extract(self, data: bytes, ctx: ExtractionContext) -> ExtractedContent:
        pages = await asyncio.to_thread(self._parse, data, ctx.document_id)
        return ExtractedContent(document_id=ctx.document_id, pages=pages)

    @staticmethod
    def _parse(data: bytes, document_id: str) -> list[Page]:
        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        pages: list[Page] = []
        try:
            for idx, ws in enumerate(wb.worksheets, start=1):
                grid: list[list[str]] = []
                for r, row in enumerate(ws.iter_rows(values_only=True)):
                    if r >= _XLSX_MAX_ROWS:
                        break
                    cells = ["" if v is None else str(v) for v in row[:_XLSX_MAX_COLS]]
                    if any(cells):
                        grid.append(cells)
                tables = (
                    [Table(page=idx, caption=ws.title, rows=grid, markdown=grid_to_markdown(grid))]
                    if grid
                    else []
                )
                heading = TextBlock(text=ws.title, role=BlockType.HEADING, page=idx, heading_level=1)
                text = ws.title + "\n" + grid_to_markdown(grid)
                pages.append(Page(number=idx, text=text, blocks=[heading], tables=tables))
        finally:
            wb.close()
        return pages or [Page(number=1, text="")]


def _heading_level_from_style(style_name: str) -> int | None:
    name = style_name.lower()
    if name == "title":
        return 1
    if name.startswith("heading"):
        digits = "".join(ch for ch in name if ch.isdigit())
        return int(digits) if digits else 1
    return None
