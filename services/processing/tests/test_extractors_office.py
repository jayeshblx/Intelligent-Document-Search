"""Tests for DOCX/PPTX/XLSX extractors using in-memory generated files."""

from __future__ import annotations

from io import BytesIO

from rag_core.models.content import BlockType
from rag_core.models.enums import DocumentType
from rag_processing.extractors.base import ExtractionContext
from rag_processing.extractors.office import DocxExtractor, PptxExtractor, XlsxExtractor


def _ctx(doc_type: DocumentType) -> ExtractionContext:
    return ExtractionContext(document_id="d1", key=f"x.{doc_type.value}", doc_type=doc_type)


def _make_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Compliance Manual", level=1)
    doc.add_paragraph("This document describes safety procedures.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Temp"
    table.cell(1, 1).text = "20C"
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Readings"
    ws.append(["Name", "Value"])
    ws.append(["x", 1])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_pptx() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    slide.shapes.title.text = "Introduction"
    slide.placeholders[1].text = "First bullet point"
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


async def test_docx_extractor_headings_paragraphs_tables() -> None:
    content = await DocxExtractor().extract(_make_docx(), _ctx(DocumentType.DOCX))
    blocks = [b for p in content.pages for b in p.blocks]
    headings = [b for b in blocks if b.role == BlockType.HEADING]
    tables = [t for p in content.pages for t in p.tables]

    assert any(h.text == "Compliance Manual" for h in headings)
    assert any("safety procedures" in b.text for b in blocks)
    assert tables and ["Item", "Value"] in tables[0].rows


async def test_xlsx_extractor_renders_sheet_as_table() -> None:
    content = await XlsxExtractor().extract(_make_xlsx(), _ctx(DocumentType.XLSX))
    assert content.page_count == 1
    page = content.pages[0]
    assert page.tables and ["Name", "Value"] in page.tables[0].rows
    assert any(b.text == "Readings" for b in page.blocks)


async def test_pptx_extractor_slide_as_page() -> None:
    content = await PptxExtractor().extract(_make_pptx(), _ctx(DocumentType.PPTX))
    assert content.page_count == 1
    blocks = content.pages[0].blocks
    assert any(b.role == BlockType.HEADING and b.text == "Introduction" for b in blocks)
    assert any("First bullet point" in b.text for b in blocks)
